import base64
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import docx
import fitz

from backend.document_parser import parse_document


class DocumentParserTests(unittest.TestCase):
    def test_pdf_extracts_text_and_renders_all_pages(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            pdf_path = Path(tmp_dir) / "scenario.pdf"
            pdf_doc = fitz.open()
            page_one = pdf_doc.new_page()
            page_one.insert_text((72, 72), "Flood map page one")
            page_two = pdf_doc.new_page()
            page_two.insert_text((72, 72), "Damage table page two")
            pdf_path.write_bytes(pdf_doc.tobytes())
            pdf_doc.close()

            with patch.dict(
                os.environ,
                {"SIMEXAI_VISION_DPI": "72", "SIMEXAI_VISION_JPEG_QUALITY": "60"},
            ):
                result = parse_document(pdf_path)

        self.assertIn("Flood map page one", result["text"])
        self.assertIn("Damage table page two", result["text"])
        self.assertEqual(result["visual_mode"], "full_page")
        self.assertEqual(result["visual_page_count"], 2)
        self.assertEqual(len(result["vision_images"]), 2)
        self.assertEqual(result["vision_images"][0]["mime_type"], "image/jpeg")
        self.assertTrue(base64.b64decode(result["vision_images"][0]["data"]).startswith(b"\xff\xd8"))

    def test_docx_extracts_text_tables_and_uses_word_pdf_rendering(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            docx_path = Path(tmp_dir) / "scenario.docx"
            document = docx.Document()
            document.add_paragraph("Cyclone scenario overview")
            table = document.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "District"
            table.rows[0].cells[1].text = "Population"
            document.save(docx_path)

            def fake_convert(_docx_path: Path, pdf_path: Path) -> None:
                pdf_doc = fitz.open()
                page = pdf_doc.new_page()
                page.insert_text((72, 72), "Rendered DOCX visual page")
                pdf_path.write_bytes(pdf_doc.tobytes())
                pdf_doc.close()

            with patch("backend.document_parser._convert_docx_to_pdf_with_word", side_effect=fake_convert):
                with patch.dict(
                    os.environ,
                    {"SIMEXAI_VISION_DPI": "72", "SIMEXAI_VISION_JPEG_QUALITY": "60"},
                ):
                    result = parse_document(docx_path)

        self.assertIn("Cyclone scenario overview", result["text"])
        self.assertIn("District | Population", result["text"])
        self.assertEqual(result["visual_mode"], "full_page")
        self.assertEqual(result["visual_page_count"], 1)
        self.assertEqual(result["vision_images"][0]["label"], "DOCX page 1")

    def test_txt_is_text_only(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            txt_path = Path(tmp_dir) / "scenario.txt"
            txt_path.write_text("Text-only scenario", encoding="utf-8")

            result = parse_document(txt_path)

        self.assertEqual(result["text"], "Text-only scenario")
        self.assertEqual(result["image_count"], 0)
        self.assertEqual(result["visual_page_count"], 0)
        self.assertEqual(result["visual_mode"], "none")
        self.assertEqual(result["vision_images"], [])

    def test_unsupported_type_raises_clear_error(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            bad_path = Path(tmp_dir) / "scenario.xlsx"
            bad_path.write_text("not supported", encoding="utf-8")

            with self.assertRaisesRegex(ValueError, "Unsupported upload type"):
                parse_document(bad_path)


if __name__ == "__main__":
    unittest.main()

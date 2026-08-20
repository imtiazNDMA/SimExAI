from __future__ import annotations

import base64
import os
import shutil
import uuid
from pathlib import Path

import docx
import fitz  # PyMuPDF


def _int_env(name: str, default: int, minimum: int, maximum: int) -> int:
    try:
        value = int(os.getenv(name, str(default)))
    except ValueError:
        return default
    return max(minimum, min(value, maximum))


def _vision_dpi() -> int:
    return _int_env("SIMEXAI_VISION_DPI", 120, 36, 300)


def _vision_jpeg_quality() -> int:
    return _int_env("SIMEXAI_VISION_JPEG_QUALITY", 70, 30, 95)


def _encode_page_as_jpeg(page: fitz.Page, dpi: int, jpeg_quality: int) -> str:
    matrix = fitz.Matrix(dpi / 72, dpi / 72)
    pixmap = page.get_pixmap(matrix=matrix, alpha=False)
    if pixmap.colorspace and pixmap.colorspace.n != 3:
        pixmap = fitz.Pixmap(fitz.csRGB, pixmap)
    image_bytes = pixmap.tobytes("jpeg", jpg_quality=jpeg_quality)
    return base64.b64encode(image_bytes).decode("ascii")


def _render_pdf_pages(pdf_path: Path, source_label: str) -> list[dict]:
    dpi = _vision_dpi()
    jpeg_quality = _vision_jpeg_quality()
    rendered_pages = []
    with fitz.open(str(pdf_path)) as pdf_doc:
        for index, page in enumerate(pdf_doc, start=1):
            rendered_pages.append(
                {
                    "label": f"{source_label} page {index}",
                    "page_number": index,
                    "mime_type": "image/jpeg",
                    "data": _encode_page_as_jpeg(page, dpi, jpeg_quality),
                }
            )
    return rendered_pages


def _extract_pdf_text_and_image_count(pdf_path: Path) -> tuple[str, int]:
    text_parts = []
    image_count = 0
    with fitz.open(str(pdf_path)) as pdf_doc:
        for page in pdf_doc:
            text_parts.append(page.get_text("text"))
            image_count += len(page.get_images(full=True))
    return "\n".join(text_parts), image_count


def _extract_docx_text_and_image_count(docx_path: Path) -> tuple[str, int]:
    document = docx.Document(docx_path)
    parts = [para.text for para in document.paragraphs if para.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    image_count = sum(
        1
        for rel in document.part.rels.values()
        if "image" in rel.reltype.lower()
    )
    return "\n".join(parts), image_count


def _convert_docx_to_pdf_with_word(docx_path: Path, pdf_path: Path) -> None:
    if os.name != "nt":
        raise RuntimeError("DOCX full-page vision rendering requires Microsoft Word on Windows.")

    try:
        import pythoncom
        import win32com.client
    except ImportError as exc:
        raise RuntimeError(
            "DOCX full-page vision rendering requires pywin32 and Microsoft Word."
        ) from exc

    word = None
    document = None
    initialized = False
    try:
        pythoncom.CoInitialize()
        initialized = True
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        document = word.Documents.Open(
            str(docx_path.resolve()),
            ConfirmConversions=False,
            ReadOnly=True,
            AddToRecentFiles=False,
        )
        document.ExportAsFixedFormat(
            OutputFileName=str(pdf_path.resolve()),
            ExportFormat=17,
            OpenAfterExport=False,
            OptimizeFor=0,
            Range=0,
            Item=0,
            IncludeDocProps=True,
            KeepIRM=True,
            CreateBookmarks=1,
            DocStructureTags=True,
            BitmapMissingFonts=True,
            UseISO19005_1=False,
        )
    except Exception as exc:
        raise RuntimeError(f"Failed to render DOCX pages with Microsoft Word: {exc}") from exc
    finally:
        if document is not None:
            try:
                document.Close(False)
            except Exception:
                pass
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass
        if initialized:
            pythoncom.CoUninitialize()


def _render_docx_pages(docx_path: Path) -> list[dict]:
    """Render DOCX pages as images via Word, or return [] if Word is unavailable.

    Page rendering needs Microsoft Word (COM, Windows only). When it isn't
    there, degrade to text-only extraction rather than failing the whole
    upload — the extracted text is still perfectly usable, it just loses
    maps, charts, and scanned content.
    """
    tmp_dir = docx_path.parent / f"simexai_docx_render_{uuid.uuid4().hex}"
    tmp_dir.mkdir(parents=True, exist_ok=False)
    try:
        pdf_path = tmp_dir / f"{docx_path.stem}.pdf"
        _convert_docx_to_pdf_with_word(docx_path, pdf_path)
        if not pdf_path.exists():
            raise RuntimeError("Microsoft Word did not create a PDF for DOCX vision rendering.")
        return _render_pdf_pages(pdf_path, "DOCX")
    except RuntimeError as exc:
        print(f"DOCX page rendering unavailable, continuing with text only: {exc}")
        return []
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def _result(text: str, image_count: int, vision_images: list[dict]) -> dict:
    return {
        "text": text.strip(),
        "image_count": image_count,
        "vision_images": vision_images,
        "visual_page_count": len(vision_images),
        "visual_mode": "full_page" if vision_images else "none",
    }


def parse_document(file_path: Path) -> dict:
    """Parse document text, media metadata, and full-page vision images.

    The uploaded file is used as LLM context only.
    """
    ext = file_path.suffix.lower()

    if ext == ".pdf":
        text, image_count = _extract_pdf_text_and_image_count(file_path)
        return _result(text, image_count, _render_pdf_pages(file_path, "PDF"))
    elif ext == ".docx":
        text, image_count = _extract_docx_text_and_image_count(file_path)
        return _result(text, image_count, _render_docx_pages(file_path))
    elif ext == ".txt":
        text = file_path.read_text(encoding="utf-8", errors="replace")
        return _result(text, 0, [])
    elif ext == ".doc":
        raise ValueError(
            "Legacy .doc files are not supported. Open the file in Word and use "
            "File > Save As to save it as .docx, then upload again."
        )
    else:
        raise ValueError(
            f"Unsupported upload type: {ext or 'unknown'}. "
            "Upload a PDF (.pdf), a Word document (.docx), or plain text (.txt)."
        )
